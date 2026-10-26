from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Dict, List, Optional

from .contact import Contact


class PhoneBook:
    def __init__(self, data_file: str | Path):
        self.data_file = Path(data_file)
        self.data_file.parent.mkdir(parents=True, exist_ok=True)
        self.contacts: List[Contact] = []
        self.load_from_json()

    def _sort_contacts(self):
        self.contacts = sorted(self.contacts, key=lambda c: c.name.lower())

    def add_contact(self, contact: Contact) -> Contact:
        if self._find_by_phone(contact.phone):
            raise ValueError("A contact with the same phone number already exists.")
        if self._find_by_email(contact.email):
            raise ValueError("A contact with the same email already exists.")

        self.contacts.append(contact)
        self._sort_contacts()
        self.save_to_json()
        return contact

    def delete_contact(self, query: str) -> Contact:
        matches = self.search(query)
        if not matches:
            raise ValueError("No contact found matching that search.")

        contact = matches[0]
        self.contacts.remove(contact)
        self.save_to_json()
        return contact

    def edit_contact(self, query: str, **updates) -> Contact:
        matches = self.search(query)
        if not matches:
            raise ValueError("No contact found matching that search.")

        contact = matches[0]

        for field, value in updates.items():
            if field == "name":
                contact.name = Contact.validate_name(value)
            elif field == "phone":
                contact.phone = Contact.normalize_phone(value)
            elif field == "email":
                contact.email = Contact.validate_email(value)
            elif field == "favorite":
                contact.favorite = bool(value)
            elif field == "notes":
                contact.notes = str(value).strip()
            else:
                raise ValueError(f"Unsupported field: {field}")

        self._sort_contacts()
        self.save_to_json()
        return contact

    def mark_favorite(self, query: str) -> Contact:
        matches = self.search(query)
        if not matches:
            raise ValueError("No contact found matching that search.")

        contact = matches[0]
        contact.favorite = True
        self.save_to_json()
        return contact

    def unmark_favorite(self, query: str) -> Contact:
        matches = self.search(query)
        if not matches:
            raise ValueError("No contact found matching that search.")

        contact = matches[0]
        contact.favorite = False
        self.save_to_json()
        return contact

    def search(self, query: str) -> List[Contact]:
        if not query or not str(query).strip():
            return []

        search_value = str(query).strip().lower()
        matches: List[Contact] = []

        try:
            normalized_phone_query = Contact.normalize_phone(search_value)
        except ValueError:
            normalized_phone_query = None

        for contact in self.contacts:
            if (
                search_value in contact.name.lower()
                or search_value in contact.email.lower()
                or (normalized_phone_query and normalized_phone_query in contact.phone.lower())
            ):
                matches.append(contact)

        return matches

    def sort_contacts(self) -> List[Contact]:
        self._sort_contacts()
        self.save_to_json()
        return self.contacts

    def get_favorites(self) -> List[Contact]:
        return [contact for contact in self.contacts if contact.favorite]

    def group_by_letter(self) -> Dict[str, List[Contact]]:
        groups: Dict[str, List[Contact]] = {}

        for contact in self.sort_contacts():
            first_letter = contact.name[0].upper() if contact.name else "#"
            if not first_letter.isalpha():
                first_letter = "#"
            groups.setdefault(first_letter, []).append(contact)

        return dict(sorted(groups.items()))

    def _find_by_phone(self, phone: str) -> Optional[Contact]:
        try:
            target = Contact.normalize_phone(phone)
        except ValueError:
            return None

        for contact in self.contacts:
            if contact.phone == target:
                return contact
        return None

    def _find_by_email(self, email: str) -> Optional[Contact]:
        try:
            target = Contact.validate_email(email)
        except ValueError:
            return None

        for contact in self.contacts:
            if contact.email == target:
                return contact
        return None

    def export_csv(self, file_path: str | Path) -> Path:
        destination = Path(file_path)
        destination.parent.mkdir(parents=True, exist_ok=True)

        fieldnames = ["id", "name", "phone", "email", "favorite", "notes"]

        with destination.open("w", newline="", encoding="utf-8") as csv_file:
            writer = csv.DictWriter(csv_file, fieldnames=fieldnames)
            writer.writeheader()
            for contact in self.contacts:
                writer.writerow(
                    {
                        "id": contact.id,
                        "name": contact.name,
                        "phone": contact.phone,
                        "email": contact.email,
                        "favorite": "Yes" if contact.favorite else "No",
                        "notes": contact.notes,
                    }
                )

        return destination

    def export_docx(self, file_path: str | Path) -> Path:
        try:
            from docx import Document
        except ImportError as exc:
            raise ImportError(
                "python-docx is required to export DOCX files. Run: pip install python-docx"
            ) from exc

        destination = Path(file_path)
        destination.parent.mkdir(parents=True, exist_ok=True)

        document = Document()
        document.add_heading("Phone Book Contacts", level=1)

        table = document.add_table(rows=1, cols=6)
        table.style = "Table Grid"
        headers = ["Name", "Phone", "Email", "Favorite", "Notes", "ID"]
        for index, header in enumerate(headers):
            table.rows[0].cells[index].text = header

        for contact in self.contacts:
            row = table.add_row().cells
            row[0].text = contact.name
            row[1].text = contact.phone
            row[2].text = contact.email
            row[3].text = "Yes" if contact.favorite else "No"
            row[4].text = contact.notes
            row[5].text = contact.id

        document.save(str(destination))
        return destination

    def save_to_json(self):
        payload = {"contacts": [contact.to_dict() for contact in self.contacts]}
        with self.data_file.open("w", encoding="utf-8") as file:
            json.dump(payload, file, indent=2, ensure_ascii=False)

    def load_from_json(self):
        if not self.data_file.exists():
            self.contacts = []
            return

        try:
            with self.data_file.open("r", encoding="utf-8") as file:
                data = json.load(file)
        except json.JSONDecodeError:
            self.contacts = []
            return

        contacts_data = data.get("contacts", [])
        self.contacts = []

        for item in contacts_data:
            try:
                self.contacts.append(Contact.from_dict(item))
            except (TypeError, ValueError):
                continue

        self._sort_contacts()
